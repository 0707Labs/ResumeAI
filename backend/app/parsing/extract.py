"""Turn an uploaded PDF, DOCX, or TXT file into ordered Lines plus format warnings."""

import io
import re
from dataclasses import dataclass, field

import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.models import ParseWarning, ParseWarningCode

from .text import SINGLE_DATE_RE, Line

MAX_BYTES = 5 * 1024 * 1024
_PRIVATE_USE = re.compile("[\\ue000-\\uf8ff]")  # Symbol/Wingdings bullet glyphs
# pdfminer renders glyphs it can't map to Unicode as "(cid:127)"; at the start of a
# line these are almost always bullet symbols.
_CID = re.compile(r"\(cid:\d+\)")
_SEGMENT_GAP_PT = 15.0  # horizontal gap that splits a PDF line into segments
_ROW_TOLERANCE_PT = 3.0


class ExtractionError(ValueError):
    """The file can't be read as a resume (empty, corrupt, scanned image)."""


class UnsupportedFileType(ExtractionError):
    pass


@dataclass
class Extracted:
    lines: list[Line]
    warnings: list[ParseWarning] = field(default_factory=list)


def extract(data: bytes, filename: str) -> Extracted:
    if len(data) > MAX_BYTES:
        raise ExtractionError("File is larger than 5 MB.")
    name = filename.lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        out = _from_pdf(data)
    elif name.endswith(".docx"):
        out = _from_docx(data)
    elif name.endswith((".txt", ".md")):
        out = _from_text(data.decode("utf-8", errors="replace"))
    else:
        raise UnsupportedFileType("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    if not any(line.text for line in out.lines):
        raise ExtractionError(
            "No text found. If this is a scanned image PDF, export a text-based PDF or DOCX "
            "instead (OCR is not supported, and most ATS systems can't read scans either)."
        )
    return out


def lines_from_text(text: str) -> list[Line]:
    return [Line(t) for t in text.splitlines() if t.strip()]


def _from_text(text: str) -> Extracted:
    return Extracted(lines_from_text(text))


# --- PDF ---------------------------------------------------------------------------


def _from_pdf(data: bytes) -> Extracted:
    lines: list[Line] = []
    warnings: list[ParseWarning] = []
    has_tables = has_images = False
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:  # pdfminer raises a zoo of exception types
        raise ExtractionError("Could not read this PDF.") from exc
    with pdf:
        for page in pdf.pages:
            words = page.extract_words(keep_blank_chars=False, use_text_flow=False)
            lines.extend(_group_words(words))
            has_images |= bool(page.images)
            has_tables |= bool(page.find_tables())
    if has_tables:
        warnings.append(_warn(ParseWarningCode.TABLES_DETECTED,
                              "Tables found. Many ATS parsers scramble table content."))
    if has_images:
        warnings.append(_warn(ParseWarningCode.IMAGES_DETECTED,
                              "Images found. ATS systems ignore text inside images."))
    if _looks_multi_column(lines):
        warnings.append(_warn(ParseWarningCode.MULTI_COLUMN_SUSPECTED,
                              "Looks like a multi-column layout. ATS parsers often read "
                              "across columns and mix up sections; a single column is safer."))
    return Extracted(lines, warnings)


def _group_words(words: list[dict]) -> list[Line]:
    rows: list[list[dict]] = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if rows and abs(rows[-1][0]["top"] - w["top"]) <= _ROW_TOLERANCE_PT:
            rows[-1].append(w)
        else:
            rows.append([w])
    lines: list[Line] = []
    for row in rows:
        row.sort(key=lambda w: w["x0"])
        segments: list[list[str]] = [[]]
        prev_x1: float | None = None
        for w in row:
            if prev_x1 is not None and w["x0"] - prev_x1 > _SEGMENT_GAP_PT:
                segments.append([])
            text = _PRIVATE_USE.sub("•", w["text"])
            text = _CID.sub("•" if prev_x1 is None else "", text)
            if text:
                segments[-1].append(text)
            prev_x1 = w["x1"]
        parts = [" ".join(s) for s in segments]
        # A lone bullet glyph followed by a gap is still one segment of text.
        if len(parts) > 1 and parts[0] in {"•", "●", "▪", "◦", "-", "–", "*", "➢", "►"}:
            parts = [f"{parts[0]} {parts[1]}", *parts[2:]]
        lines.append(Line(" ".join(parts), segments=parts, x0=row[0]["x0"]))
    return lines


def _looks_multi_column(lines: list[Line]) -> bool:
    # Right-aligned dates ("Acme Corp      Jan 2020 – Present") are normal; a wide
    # gap with prose on both sides, on many lines, means two columns side by side.
    if len(lines) < 8:
        return False
    split = 0
    for line in lines:
        if len(line.segments) < 2:
            continue
        right = " ".join(line.segments[1:])
        if not SINGLE_DATE_RE.search(right) and len(right.split()) >= 2:
            split += 1
    return split / len(lines) >= 0.3


# --- DOCX ----------------------------------------------------------------------------


def _from_docx(data: bytes) -> Extracted:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Could not read this DOCX file.") from exc
    lines: list[Line] = []
    warnings: list[ParseWarning] = []
    has_tables = False
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            line = _docx_paragraph(Paragraph(child, doc))
            if line:
                lines.append(line)
        elif child.tag == qn("w:tbl"):
            has_tables = True
            lines.extend(_docx_table(Table(child, doc)))
    if has_tables:
        warnings.append(_warn(ParseWarningCode.TABLES_DETECTED,
                              "Tables found. Many ATS parsers scramble table content."))
    if doc.inline_shapes:
        warnings.append(_warn(ParseWarningCode.IMAGES_DETECTED,
                              "Images found. ATS systems ignore text inside images."))
    return Extracted(lines, warnings)


def _docx_paragraph(p: Paragraph) -> Line | None:
    text = p.text.strip()
    if not text:
        return None
    style = (p.style.name if p.style is not None else "") or ""
    num_pr = p._p.pPr is not None and p._p.pPr.numPr is not None
    is_list = num_pr or "List" in style
    # Tabs often separate "Acme Corp<TAB>Jan 2020 – Present".
    segments = [s for s in re.split(r"\t+", text) if s.strip()]
    return Line(text, segments=segments, is_list_item=is_list)


def _docx_table(table: Table) -> list[Line]:
    lines: list[Line] = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            if id(cell._tc) in seen:  # merged cells repeat
                continue
            seen.add(id(cell._tc))
            cells.append(cell)
        if all(len([p for p in c.paragraphs if p.text.strip()]) <= 1 for c in cells):
            texts = [c.text.strip() for c in cells if c.text.strip()]
            if texts:
                lines.append(Line(" ".join(texts), segments=texts))
        else:
            for c in cells:
                for p in c.paragraphs:
                    line = _docx_paragraph(p)
                    if line:
                        lines.append(line)
    return lines


def _warn(code: ParseWarningCode, message: str) -> ParseWarning:
    return ParseWarning(code=code, message=message)
