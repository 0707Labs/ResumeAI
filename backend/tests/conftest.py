"""Build PDF/DOCX resumes at test time so no binary fixtures live in git.

Database tests get a throwaway Postgres database per run (migrated with Alembic)
and a per-test transaction that is rolled back."""

import io
import os
import uuid
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas
from fastapi import Depends, Request, Response
from fastapi.testclient import TestClient
from pydantic import BaseModel
from sqlalchemy import create_engine, make_url, text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import NullPool

from app.api.signin import AuthStep, sign_in_or_gate
from app.auth.invites import generate_code, hash_code
from app.auth.mailer import Mail, set_mailer
from app.auth.ratelimit import limiter
from app.auth.sessions import start_session
from app.db.models import User
from app.db.session import get_db
from app.main import create_app

FIXTURES = Path(__file__).parent / "fixtures"
BACKEND = Path(__file__).parent.parent
# A server we may create and drop databases on; the local socket by default.
ADMIN_URL = os.environ.get("RT_TEST_ADMIN_DATABASE_URL", "postgresql+psycopg:///postgres")
ORIGIN = "http://localhost:3000"  # the default RT_PUBLIC_URL
TABLES = ("users, accounts, auth_sessions, pending_auth, email_tokens, invites, invite_redemptions,"
          " waitlist, consent_records, consent_notices, credit_transactions, passes,"
          " payments, billing_customers, processed_webhook_events, subscriptions, billing_reviews, llm_usage, analyses, downloads, referrals, app_settings,"
          " job_preferences, resumes, resume_versions")


def fixture_text(name: str) -> str:
    return (FIXTURES / name).read_text()


def make_pdf(rows: list[tuple], *, blank: bool = False) -> bytes:
    """rows: (x, text) for left text, or (x, text, right_text) for a right-aligned
    date on the same line. One row per line, top to bottom."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=LETTER)
    width, height = LETTER
    if not blank:
        y = height - 60
        c.setFont("Helvetica", 10)
        for row in rows:
            c.drawString(row[0], y, row[1])
            if len(row) > 2:
                c.drawRightString(width - 50, y, row[2])
            y -= 14
    else:
        c.rect(50, 50, 100, 100)  # a drawing with no text, like a scanned page
    c.showPage()
    c.save()
    return buf.getvalue()


@pytest.fixture
def analyst_pdf() -> bytes:
    L, B, W = 50, 60, 68  # left margin, bullet glyph x, wrapped-text x
    return make_pdf([
        (L, "Ravi Menon"),
        (L, "ravi.menon@example.com | +1 415 555 0134 | github.com/ravimenon"),
        (L, "EXPERIENCE"),
        (L, "Software Engineer | Globex Corporation", "Mar 2021 – Present"),
        (B, "• Contributed to migrating billing services from a monolith to"),
        (W, "microservices in Java and Spring Boot"),
        (B, "• Reduced p95 API latency by 40% by adding Redis caching"),
        (L, "Software Engineering Intern | Initech", "Jun 2020 – Aug 2020"),
        (B, "• Wrote unit tests for the payments module with JUnit"),
        (L, "EDUCATION"),
        (L, "B.S. Computer Science, University of Washington", "2016 – 2020"),
        (L, "SKILLS"),
        (L, "Java, Python, SQL, Docker, Redis"),
    ])


@pytest.fixture
def two_column_pdf() -> bytes:
    rows = [(50, "Priya Shah")]
    for i in range(10):
        rows.append((50, f"Left column sentence number {i} here", f"right column words {i} here"))
    return make_pdf(rows)


@pytest.fixture
def blank_pdf() -> bytes:
    return make_pdf([], blank=True)


@pytest.fixture
def analyst_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Meera Iyer", style="Title")
    doc.add_paragraph("meera.iyer@example.com | +44 20 7946 0958")
    doc.add_paragraph("Professional Experience", style="Heading 1")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Marketing Analyst, Contoso Ltd"
    table.rows[0].cells[1].text = "Feb 2019 – Present"
    doc.add_paragraph("Helped run weekly campaign reporting in Excel", style="List Bullet")
    doc.add_paragraph("Grew newsletter signups 25% through A/B tests on subject lines",
                      style="List Bullet")
    doc.add_paragraph("Product Analyst\tJan 2017 – Jan 2019")
    doc.add_paragraph("Fabrikam Inc")
    doc.add_paragraph("Assisted with SQL queries for the quarterly churn review",
                      style="List Bullet")
    doc.add_paragraph("Education", style="Heading 1")
    doc.add_paragraph("MBA, London Business School, 2016")
    doc.add_paragraph("Skills", style="Heading 1")
    doc.add_paragraph("Excel; SQL; Google Analytics; Tableau")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def alembic_config(url: str) -> Config:
    cfg = Config(str(BACKEND / "alembic.ini"))
    cfg.set_main_option("sqlalchemy.url", url.replace("%", "%%"))
    return cfg


@pytest.fixture(scope="session")
def pg_url() -> str:
    """URL of a fresh database migrated to head; dropped after the run.
    Skips DB tests when no Postgres is reachable, except in CI where that fails."""
    admin = create_engine(ADMIN_URL, isolation_level="AUTOCOMMIT")
    name = f"rt_test_{uuid.uuid4().hex[:8]}"
    try:
        with admin.connect() as conn:
            conn.execute(text(f'CREATE DATABASE "{name}"'))
    except Exception as exc:  # noqa: BLE001 - any connection failure means "no Postgres here"
        admin.dispose()
        if os.environ.get("CI"):
            raise
        pytest.skip(f"Postgres not reachable at {ADMIN_URL}: {exc.__class__.__name__}")
    url = make_url(ADMIN_URL).set(database=name).render_as_string(hide_password=False)
    try:
        command.upgrade(alembic_config(url), "head")
        yield url
    finally:
        with admin.connect() as conn:
            conn.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
        admin.dispose()


@pytest.fixture(scope="session")
def pg_engine(pg_url):
    engine = create_engine(pg_url)
    yield engine
    engine.dispose()


@pytest.fixture
def db(pg_engine):
    """A Session inside one outer transaction that is rolled back after the test.
    Commits inside the test become savepoints."""
    with pg_engine.connect() as conn:
        outer = conn.begin()
        session = Session(bind=conn, join_transaction_mode="create_savepoint")
        try:
            yield session
        finally:
            session.close()
            outer.rollback()


class Outbox(list):
    """Records magic-link mail instead of sending it."""

    def send(self, mail: Mail) -> None:
        self.append(mail)

    def last_token(self) -> str:
        return self[-1].link.split("#token=", 1)[1]


class IdentityIn(BaseModel):
    provider: str
    provider_account_id: str
    verified_email: str | None = None


@pytest.fixture
def outbox():
    box = Outbox()
    set_mailer(box)
    yield box
    set_mailer(None)


@pytest.fixture
def api_app(pg_url, pg_engine, outbox):
    """The real app on the test database, plus test-only sign-in shortcuts:
    /test/login/{user_id} and /test/identity (what an OAuth callback will do).
    Tables are truncated afterwards."""
    engine = create_async_engine(pg_url, poolclass=NullPool)
    sessions = async_sessionmaker(engine, expire_on_commit=False)

    async def test_db():
        async with sessions() as s:
            yield s

    app = create_app()
    app.dependency_overrides[get_db] = test_db

    @app.post("/test/login/{user_id}")
    async def test_login(user_id: uuid.UUID, request: Request, response: Response,
                         db: AsyncSession = Depends(get_db)) -> None:
        await start_session(db, request, response, await db.get(User, user_id), provider="google")

    @app.post("/test/identity")
    async def test_identity(body: IdentityIn, request: Request, response: Response,
                            db: AsyncSession = Depends(get_db)) -> AuthStep:
        return await sign_in_or_gate(db, request, response, provider=body.provider,
                                      provider_account_id=body.provider_account_id,
                                      verified_email=body.verified_email)

    limiter.reset()
    yield app
    limiter.reset()
    with pg_engine.begin() as conn:
        conn.execute(text("SET LOCAL lock_timeout = '5s'"))  # a leaked connection fails, not hangs
        conn.execute(text(f"TRUNCATE {TABLES} CASCADE"))


@pytest.fixture
def new_browser(api_app):
    """Factory for independent cookie jars against the same app."""
    clients = []

    def make() -> TestClient:
        c = TestClient(api_app, headers={"Origin": ORIGIN})
        c.__enter__()
        clients.append(c)
        return c

    yield make
    for c in clients:
        c.__exit__(None, None, None)


@pytest.fixture
def client(new_browser) -> TestClient:
    return new_browser()


@pytest.fixture
def user_id(pg_engine) -> uuid.UUID:
    """An active user, ada@example.com, with no login methods yet."""
    with pg_engine.begin() as conn:
        return conn.execute(text(
            "INSERT INTO users (email, email_verified_at, activated_at, age_confirmed_at,"
            " terms_version, terms_accepted_at) VALUES ('ada@example.com', now(), now(), now(),"
            " 'v1', now()) RETURNING id"
        )).scalar_one()


@pytest.fixture
def notices(pg_engine) -> dict[str, str]:
    ids = {}
    with pg_engine.begin() as conn:
        for purpose in ("service", "age_18_plus", "research", "waitlist"):
            ids[purpose] = str(conn.execute(text(
                "INSERT INTO consent_notices (purpose, version, body_sha256, body_url, published_at)"
                " VALUES (:p, '2026-11-v1', :h, :u, now() - interval '1 minute') RETURNING id"
            ), {"p": purpose, "h": "0" * 64, "u": f"/legal/{purpose}"}).scalar_one())
    return ids


@pytest.fixture
def make_invite(pg_engine):
    def make(email: str | None = None, max_uses: int = 1, extra: str = "") -> str:
        code = generate_code()
        with pg_engine.begin() as conn:
            conn.execute(text(
                "INSERT INTO invites (code_hash, code_hint, email, max_uses"
                + (", " + extra.split("=")[0] if extra else "") + ") VALUES (:h, :hint, :e, :m"
                + (", " + extra.split("=", 1)[1] if extra else "") + ")"
            ), {"h": hash_code(code), "hint": code[-4:], "e": email, "m": max_uses})
        return code
    return make
