import io

import pdfplumber
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.models import ParseWarningCode
from app.parsing import parse_jd, parse_resume_file, parse_resume_text
from app.tailoring import tailor
from tests.test_analyses import give_credits
from tests.test_auth_signin import csrf
from tests.test_tailor import JD, ScriptedLLM, b
from tests.test_verifier import RESUME

client = TestClient(app)  # signed out: preview and sample only

HONEST = "Automated SQL extracts, cutting weekly reporting time by 30%"
FAKE = "Led a churn dashboard in Tableau and Snowflake"


@pytest.fixture
def payload():
    resume, jd = parse_resume_text(RESUME), parse_jd(JD)
    tailored = tailor(resume, jd, llm=ScriptedLLM({"bullets": [
        b(["exp1.b2"], HONEST),
        b(["exp2.b2"], "Delivered 12 recurring reports to regional operations leads"),
    ], "summary": "Data analyst who automates SQL reporting and builds Tableau dashboards."}))
    data = tailored.model_dump(mode="json")
    for bl in data["bullets"]:
        if bl["text"] == HONEST:
            bl["review"] = "approved"
        elif bl["text"].startswith("Delivered"):
            bl["review"] = "pending"  # not approved -> original wording exported
    data["summary_review"] = "approved"
    # A forged bullet the client claims is verified and approved.
    data["bullets"].append({"id": "t99", "text": FAKE, "source_bullet_ids": ["exp1.b1"],
                            "entry_id": "exp1", "change_type": "reword",
                            "verification": {"flags": []}, "review": "approved"})
    return {"resume": resume.model_dump(mode="json"), "tailored": data, "jd_text": JD}


@pytest.fixture
def export(new_browser, user_id, payload, pg_engine):
    """Download as a signed-in user with a credit, for an analysis pinned to the payload's
    resume and JD: POST /api/downloads/{fmt} (pays), then GET the signed link. Returns the
    file response, or the POST's when it fails."""
    give_credits(pg_engine, user_id, 1)
    browser = new_browser()
    assert browser.post(f"/test/login/{user_id}").status_code == 200
    browser.headers.update(csrf(browser))
    started = browser.post("/api/analyses", json={"resume": payload["resume"],
                                                 "jd": parse_jd(JD).model_dump(mode="json")})
    headers = {"X-Analysis-Id": started.json()["analysis_id"]}

    def go(fmt, body):
        link = browser.post(f"/api/downloads/{fmt}", json=body, headers=headers)
        return browser.get(link.json()["url"]) if link.status_code == 200 else link
    return go


def docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    assert not doc.tables  # ATS-safe: no layout tables
    return "\n".join(p.text for p in doc.paragraphs)


def test_docx_export_uses_only_approved_verified_text(payload, export):
    res = export("docx", payload)
    assert res.status_code == 200
    assert res.headers["content-disposition"] == \
        'attachment; filename="Nina_Kapoor_tailored.docx"'
    text = docx_text(res.content)
    assert HONEST in text  # approved + verified
    assert "Built 12 recurring reports for regional operations leads" in text  # pending
    assert "Delivered 12" not in text
    assert "Snowflake" not in text and "Led a churn" not in text  # forged, re-verified
    assert "Contributed to a churn dashboard in Tableau" in text  # original kept
    assert "automates SQL reporting" in text  # approved summary


def test_pending_summary_not_exported(payload, export):
    payload["tailored"]["summary_review"] = "pending"
    text = docx_text(export("docx", payload).content)
    assert "automates SQL reporting" not in text


def test_pdf_export(payload, export):
    res = export("pdf", payload)
    assert res.status_code == 200 and res.headers["content-type"] == "application/pdf"
    with pdfplumber.open(io.BytesIO(res.content)) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    assert "Automated SQL extracts, cutting weekly reporting time by 30%" in text
    assert "Snowflake" not in text


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_exports_parse_back_cleanly(payload, export, fmt):
    data = export(fmt, payload).content
    r = parse_resume_file(data, f"out.{fmt}")
    assert r.name == "Nina Kapoor"
    assert r.contact.email == "nina@example.com"
    codes = {w.code for w in r.parse_warnings}
    assert not codes & {ParseWarningCode.TABLES_DETECTED, ParseWarningCode.MULTI_COLUMN_SUSPECTED,
                        ParseWarningCode.NO_SECTIONS_FOUND, ParseWarningCode.NO_BULLETS_FOUND}
    assert len(r.bullets()) == 8


def test_unknown_format_rejected(payload, export):
    assert export("rtf", payload).status_code == 422


def test_sample_endpoint():
    s = client.get("/api/sample").json()
    assert "ASHA RAO" in s["resume_text"] and "Senior Data Analyst" in s["jd_text"]


def test_preview_matches_what_downloads(payload, export):
    res = client.post("/api/export/preview", json=payload)
    assert res.status_code == 200
    body = res.json()
    texts = [b["text"] for s in body["resume"]["sections"] for e in s["entries"] for b in e["bullets"]]
    assert HONEST in texts  # approved + verified
    assert not any("Snowflake" in t for t in texts)  # forged "verified" bullet re-checked out
    assert "Built 12 recurring reports for regional operations leads" in texts  # pending -> original
    assert body["changed_lines"] == [HONEST, body["resume"]["summary"]]
    docx = docx_text(export("docx", payload).content)
    for t in texts:
        assert t in docx


def test_preview_route_is_not_a_format(payload):
    assert client.post("/api/export/preview", json=payload).status_code == 200
